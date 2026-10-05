#!/usr/bin/env python3
"""Build test/test-deck.pptx: 4 visible slides + 1 hidden slide.

  1  multi-sentence notes, with [pause] and [click] stage directions
  2  (hidden) has notes that must NOT be spoken or shown
  3  multi-sentence notes, acronyms for the pronunciation file, and an inline
     [word](/phonemes/) override that must be kept, not stripped as a stage direction
  4  no notes at all
  5  one long paragraph (well past one TTS chunk)
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt

SLIDES = [
    ("Slide 1: Welcome", False,
     "Welcome to the deck narrator test. [pause] This first slide has a few sentences of notes. "
     "The bracketed stage directions should never be spoken aloud. [click] That is all for slide one."),
    ("HIDDEN SLIDE", True,
     "This slide is hidden. If you can hear this sentence, hidden slide handling is broken."),
    ("Slide 3: Acronyms", False,
     "NGINX runs on an NVMe drive. The GPU is an RTX card! Does the pronunciation file work? "
     "We [present](/pɹˈɛzənt/) the present, on purpose backwards."),
    ("Slide 4: No notes", False, ""),
    ("Slide 5: The long one", False,
     "This slide carries one long paragraph of speaker notes, written to be well past the length of a "
     "single text to speech chunk, so that the tool has to split it at sentence boundaries and join "
     "the audio back together without dropping anything. The first sentence sets up the idea. The "
     "second sentence keeps going, adding a clause, and then another clause, and then one more, just "
     "to make it a little longer than strictly necessary. Here is a third sentence about photographs, "
     "video libraries, and the machines that store them in a small office. A fourth sentence talks about "
     "backups, because there is no such thing as having too many of them. The fifth sentence notes "
     "that animations are flattened, which is expected. The sixth sentence is here only for length. "
     "And finally, the last sentence ends the paragraph, so if you hear the word pineapple, nothing "
     "was truncated. Pineapple."),
]

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)   # 16:9
for title, hidden, notes in SLIDES:
    s = prs.slides.add_slide(prs.slide_layouts[5])                 # title only
    s.shapes.title.text = title
    s.shapes.title.text_frame.paragraphs[0].runs[0].font.size = Pt(54)
    if hidden:
        s._element.set("show", "0")
    if notes:
        s.notes_slide.notes_text_frame.text = notes
out = Path(__file__).with_name("test-deck.pptx")
prs.save(out)
print(out)
