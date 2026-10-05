#!/usr/bin/env python3
"""Build test/anim-deck.pptx: animation builds and voice switching.

The animation XML is written the way PowerPoint writes it: one outer <p:par>
per click in the main sequence, "with/after previous" effects nested inside
the click that starts them, each effect a <p:cTn presetClass nodeType> with
<p:set style.visibility> plus an <p:animEffect> behaviour, and <p:bldP> build
entries for paragraph builds.

  1  bullets built one paragraph per click; notes click three times
  2  shapes: an auto-start subtitle, click 1 brings two boxes in together
     (click + with previous), click 2 removes one, click 3 is emphasis only.
     The notes click only twice, so click 3 plays after the narration (warning)
  3  two voices in conversation; ends on am_michael
  4  no voice marker, so am_michael carries over from slide 3
  5  no notes, two builds: they play through, then rest on the final state
"""
from pathlib import Path
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NSDECL = f'xmlns:p="{P}" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'


class Timing:
    """Writes PowerPoint-style <p:timing> for a list of clicks."""

    def __init__(self):
        self.next_id = 3            # 1 = tmRoot, 2 = mainSeq
        self.builds = []

    def nid(self):
        self.next_id += 1
        return self.next_id - 1

    def target(self, spid, para):
        if para is None:
            return f'<p:spTgt spid="{spid}"/>'
        return f'<p:spTgt spid="{spid}"><p:txEl><p:pRg st="{para}" end="{para}"/></p:txEl></p:spTgt>'

    def effect(self, cls, node_type, spid, para=None, delay=0):
        tgt = self.target(spid, para)
        if para is not None and (spid, "p") not in self.builds:
            self.builds.append((spid, "p"))
        grp = 0
        if cls == "emph":
            body = (f'<p:animScale><p:cBhvr><p:cTn id="{self.nid()}" dur="250" autoRev="1" fill="hold"/>'
                    f'<p:tgtEl>{tgt}</p:tgtEl></p:cBhvr><p:by x="105000" y="105000"/></p:animScale>')
            preset = 'presetID="32" presetClass="emph" presetSubtype="0"'
        else:
            vis = "visible" if cls == "entr" else "hidden"
            trans = "in" if cls == "entr" else "out"
            fade = (f'<p:animEffect transition="{trans}" filter="fade"><p:cBhvr><p:cTn id="{self.nid()}" dur="500"/>'
                    f'<p:tgtEl>{tgt}</p:tgtEl></p:cBhvr></p:animEffect>')
            setv = (f'<p:set><p:cBhvr><p:cTn id="{self.nid()}" dur="1" fill="hold"><p:stCondLst><p:cond delay="{0 if cls == "entr" else 499}"/>'
                    f'</p:stCondLst></p:cTn><p:tgtEl>{tgt}</p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName>'
                    f'</p:attrNameLst></p:cBhvr><p:to><p:strVal val="{vis}"/></p:to></p:set>')
            body = setv + fade if cls == "entr" else fade + setv
            preset = f'presetID="10" presetClass="{cls}" presetSubtype="0"'
        return (f'<p:par><p:cTn id="{self.nid()}" {preset} fill="hold" grpId="{grp}" nodeType="{node_type}">'
                f'<p:stCondLst><p:cond delay="{delay}"/></p:stCondLst><p:childTnLst>{body}</p:childTnLst></p:cTn></p:par>')

    def click_group(self, subgroups, auto=False):
        """subgroups: lists of effect XML; the first subgroup starts on the click
        (or automatically), each later one is an "after previous"."""
        inner = ""
        delay = 0
        for effects in subgroups:
            inner += (f'<p:par><p:cTn id="{self.nid()}" fill="hold"><p:stCondLst><p:cond delay="{delay}"/></p:stCondLst>'
                      f'<p:childTnLst>{"".join(effects)}</p:childTnLst></p:cTn></p:par>')
            delay += 500
        cond = ('<p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond>'
                if auto else '<p:cond delay="indefinite"/>')
        return (f'<p:par><p:cTn id="{self.nid()}" fill="hold"><p:stCondLst>{cond}</p:stCondLst>'
                f'<p:childTnLst>{inner}</p:childTnLst></p:cTn></p:par>')

    def xml(self, groups):
        bld = "".join(f'<p:bldP spid="{spid}" grpId="0" build="p"/>' for spid, _ in self.builds)
        return etree.fromstring(
            f'<p:timing {NSDECL}><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">'
            f'<p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq">'
            f'<p:childTnLst>{"".join(groups)}</p:childTnLst></p:cTn>'
            f'<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
            f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>'
            f'</p:seq></p:childTnLst></p:cTn></p:par></p:tnLst>'
            + (f'<p:bldLst>{bld}</p:bldLst>' if bld else '') + '</p:timing>')


def attach(slide, timing_el):
    sld = slide._element
    # <p:timing> goes after <p:clrMapOvr>/<p:transition>, before <p:extLst>.
    ext = sld.find(f"{{{P}}}extLst")
    if ext is not None:
        ext.addprevious(timing_el)
    else:
        sld.append(timing_el)


def box(slide, text, x, y, color):
    shp = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(3.2), Inches(1.6))
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGBColor.from_string(color)
    shp.text_frame.text = text
    shp.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    return shp


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    # 1: bullets by paragraph
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Builds: bullets"
    body = s.placeholders[1]
    body.text_frame.text = "First bullet"
    for t in ("Second bullet", "Third bullet"):
        body.text_frame.add_paragraph().text = t
    t = Timing()
    attach(s, t.xml([t.click_group([[t.effect("entr", "clickEffect", body.shape_id, i)]]) for i in range(3)]))
    s.notes_slide.notes_text_frame.text = (
        "Here is a slide that builds. [click] First, the first bullet appears. "
        "[click] Then the second one. [click] And finally the third.")

    # 2: shapes, auto-start, with-previous, exit, emphasis
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Builds: shapes"
    sub = s.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(8), Inches(0.8))
    sub.text_frame.text = "Subtitle that fades in on its own"
    a = box(s, "Box A (leaves on click 2)", 0.8, 3.2, "4472C4")
    b = box(s, "Box B (click 1)", 4.9, 3.2, "70AD47")
    c = box(s, "Box C (with B)", 9.0, 3.2, "ED7D31")
    t = Timing()
    attach(s, t.xml([
        t.click_group([[t.effect("entr", "afterEffect", sub.shape_id)]], auto=True),
        t.click_group([[t.effect("entr", "clickEffect", b.shape_id), t.effect("entr", "withEffect", c.shape_id)]]),
        t.click_group([[t.effect("exit", "clickEffect", a.shape_id)]]),
        t.click_group([[t.effect("emph", "clickEffect", b.shape_id)]]),
    ]))
    s.notes_slide.notes_text_frame.text = (
        "Box A starts on screen, and the subtitle fades in by itself. [click] Boxes B and C arrive together. "
        "[click] Now box A is gone.")

    # 3: voices in conversation
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Voices: a conversation"
    s.notes_slide.notes_text_frame.text = (
        "This is the default voice, opening the show. [voice: am_michael] And I'm the second voice, an American man. "
        "[voice: bf_emma] I'm British, and slightly louder than you'd think. [voice: default] Back to the default. "
        "[voice: am_michael] Michael gets the last word on this slide.")

    # 4: voice carries over
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Voices: carry-over"
    s.notes_slide.notes_text_frame.text = "Still Michael here, because no marker switched me back."

    # 5: no notes, two builds
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Builds without notes"
    d = box(s, "Appears", 1.5, 3.0, "7030A0")
    e = box(s, "Then this", 6.0, 3.0, "C00000")
    t = Timing()
    attach(s, t.xml([t.click_group([[t.effect("entr", "clickEffect", d.shape_id)]]),
                     t.click_group([[t.effect("entr", "clickEffect", e.shape_id)]])]))

    out = Path(__file__).with_name("anim-deck.pptx")
    prs.save(out)
    print(out)


if __name__ == "__main__":
    main()
