#!/usr/bin/env python3
"""Unit tests for builds.py (animation parsing and state rendering) and for the
[click] / [voice: ...] script in narrate.py. No GPU, no LibreOffice.

Run:  .venv/bin/python -m unittest discover -s test -v
"""
import sys
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))
import builds  # noqa: E402
import narrate  # noqa: E402
from make_anim_deck import Timing, attach, box  # noqa: E402
from pptx import Presentation  # noqa: E402

A = builds.A


def new_slide(layout=5):
    prs = Presentation()
    return prs, prs.slides.add_slide(prs.slide_layouts[layout])


def bullets(slide, n=3):
    body = slide.placeholders[1]
    body.text_frame.text = "one"
    for i in range(1, n):
        body.text_frame.add_paragraph().text = f"p{i}"
    return body


class ParseBuilds(unittest.TestCase):
    def test_no_timing(self):
        _, s = new_slide()
        b = builds.parse_builds(s._element)
        self.assertEqual(b.states, [frozenset()])
        self.assertEqual(b.clicks, 0)

    def test_paragraph_build(self):
        _, s = new_slide(1)
        body = bullets(s)
        t = Timing()
        attach(s, t.xml([t.click_group([[t.effect("entr", "clickEffect", body.shape_id, i)]]) for i in range(3)]))
        b = builds.parse_builds(s._element)
        sid = str(body.shape_id)
        self.assertEqual(b.states, [frozenset({(sid, 0), (sid, 1), (sid, 2)}),
                                    frozenset({(sid, 1), (sid, 2)}), frozenset({(sid, 2)}), frozenset()])

    def test_with_previous_joins_the_click(self):
        _, s = new_slide()
        x, y = box(s, "x", 1, 1, "000000"), box(s, "y", 5, 1, "000000")
        t = Timing()
        attach(s, t.xml([t.click_group([[t.effect("entr", "clickEffect", x.shape_id),
                                         t.effect("entr", "withEffect", y.shape_id)]])]))
        b = builds.parse_builds(s._element)
        self.assertEqual(b.clicks, 1)
        self.assertEqual(b.states[1], frozenset())

    def test_after_previous_inside_a_click_joins_it(self):
        _, s = new_slide()
        x, y = box(s, "x", 1, 1, "000000"), box(s, "y", 5, 1, "000000")
        t = Timing()
        attach(s, t.xml([t.click_group([[t.effect("entr", "clickEffect", x.shape_id)],
                                        [t.effect("entr", "afterEffect", y.shape_id)]])]))
        self.assertEqual(builds.parse_builds(s._element).clicks, 1)

    def test_auto_start_is_state_zero(self):
        _, s = new_slide()
        x, y = box(s, "x", 1, 1, "000000"), box(s, "y", 5, 1, "000000")
        t = Timing()
        attach(s, t.xml([t.click_group([[t.effect("entr", "afterEffect", x.shape_id)]], auto=True),
                         t.click_group([[t.effect("entr", "clickEffect", y.shape_id)]])]))
        b = builds.parse_builds(s._element)
        self.assertEqual(b.clicks, 1)
        self.assertEqual(b.states[0], frozenset({(str(y.shape_id), None)}))   # x already in
        self.assertEqual(b.auto_start, 1)

    def test_exit_and_emphasis(self):
        _, s = new_slide()
        x = box(s, "x", 1, 1, "000000")
        t = Timing()
        attach(s, t.xml([t.click_group([[t.effect("emph", "clickEffect", x.shape_id)]]),
                         t.click_group([[t.effect("exit", "clickEffect", x.shape_id)]])]))
        b = builds.parse_builds(s._element)
        xid = (str(x.shape_id), None)
        self.assertEqual(b.states, [frozenset(), frozenset(), frozenset({xid})])  # emphasis: a click, no change
        self.assertEqual(b.ignored, {"emph"})

    def test_enter_then_exit(self):
        _, s = new_slide()
        x = box(s, "x", 1, 1, "000000")
        t = Timing()
        attach(s, t.xml([t.click_group([[t.effect("entr", "clickEffect", x.shape_id)]]),
                         t.click_group([[t.effect("exit", "clickEffect", x.shape_id)]])]))
        xid = (str(x.shape_id), None)
        self.assertEqual(builds.parse_builds(s._element).states, [frozenset({xid}), frozenset(), frozenset({xid})])


class ApplyState(unittest.TestCase):
    def test_removes_shape_and_timing(self):
        _, s = new_slide()
        x, y = box(s, "x", 1, 1, "000000"), box(s, "y", 5, 1, "000000")
        t = Timing()
        attach(s, t.xml([t.click_group([[t.effect("entr", "clickEffect", x.shape_id)]])]))
        missing = builds.apply_state(s._element, {(str(x.shape_id), None)})
        self.assertEqual(missing, [])
        ids = [sh.shape_id for sh in s.shapes]
        self.assertNotIn(x.shape_id, ids)
        self.assertIn(y.shape_id, ids)
        self.assertIsNone(s._element.find(f"{{{builds.P}}}timing"))

    def test_hides_paragraph_keeps_text(self):
        _, s = new_slide(1)
        body = bullets(s)
        builds.apply_state(s._element, {(str(body.shape_id), 1)})
        paras = body.text_frame._txBody.findall(f"{{{A}}}p")
        self.assertEqual([p.xpath("string(.)") for p in paras], ["one", "p1", "p2"])   # text kept: layout unchanged
        hidden = paras[1]
        self.assertIsNotNone(hidden.find(f"{{{A}}}pPr/{{{A}}}buNone"))
        alpha = hidden.find(f"{{{A}}}r/{{{A}}}rPr/{{{A}}}solidFill/{{{A}}}srgbClr/{{{A}}}alpha")
        self.assertEqual(alpha.get("val"), "0")
        self.assertIsNone(paras[0].find(f"{{{A}}}pPr/{{{A}}}buNone"))

    def test_missing_target_reported(self):
        _, s = new_slide()
        self.assertEqual(builds.apply_state(s._element, {("999", None)}), [("999", None)])


class Script(unittest.TestCase):
    def parse(self, raw, voice="af_heart"):
        return narrate.parse_script(raw, voice, "af_heart", 1)

    def test_clicks_split_steps(self):
        parts, clicks, _ = self.parse("Intro. [click] One. [CLICK] Two.")
        self.assertEqual(clicks, 2)
        self.assertEqual([(p["step"], p["text"]) for p in parts], [(0, "Intro."), (1, "One."), (2, "Two.")])

    def test_voice_switch_and_default(self):
        parts, _, end = self.parse("A. [voice: am_michael] B. [Voice:bf_emma] C. [voice: default] D.")
        self.assertEqual([p["voice"] for p in parts], ["af_heart", "am_michael", "bf_emma", "af_heart"])
        self.assertEqual(end, "af_heart")

    def test_voice_carries_out(self):
        _, _, end = self.parse("A. [voice: am_michael] B.")
        self.assertEqual(end, "am_michael")
        parts, _, _ = narrate.parse_script("Still me.", end, "af_heart", 2)
        self.assertEqual(parts[0]["voice"], "am_michael")

    def test_bad_voice_names_slide(self):
        with self.assertRaisesRegex(narrate.Fail, "slide 1: unknown voice 'bob'"):
            self.parse("[voice: bob] hi")

    def test_punctuation_after_click_stays_with_words(self):
        parts, clicks, _ = self.parse("Done [click].")
        self.assertEqual(clicks, 1)
        self.assertEqual(parts, [{"step": 0, "voice": "af_heart", "text": "Done."}])

    def test_other_directions_still_stripped(self):
        parts, _, _ = self.parse("Hi [pause] there [laughs].")
        self.assertEqual(parts[0]["text"], "Hi there.")

    def test_override_is_not_a_marker(self):
        parts, clicks, _ = self.parse("We [present](/pɹizˈɛnt/) it.")
        self.assertEqual(clicks, 0)
        self.assertEqual(parts[0]["text"], "We [present](/pɹizˈɛnt/) it.")

    def test_fold_without_builds(self):
        parts, _, _ = self.parse("A. [click] B. [click] C.")
        self.assertEqual(narrate.fold_parts(parts, 1), [{"step": 0, "voice": "af_heart", "text": "A. B. C."}])

    def test_fold_clamps_extra_clicks(self):
        parts, _, _ = self.parse("A. [click] B. [click] C.")
        self.assertEqual([(p["step"], p["text"]) for p in narrate.fold_parts(parts, 2)], [(0, "A."), (1, "B. C.")])

    def test_fold_keeps_voice_boundaries(self):
        parts, _, _ = self.parse("A. [voice: am_michael] B.")
        self.assertEqual(len(narrate.fold_parts(parts, 1)), 2)


class Plan(unittest.TestCase):
    def slide(self, raw, n_states):
        parts, clicks, _ = narrate.parse_script(raw, "af_heart", "af_heart", 1)
        b = builds.Builds([frozenset()] * n_states, set(), 0)
        return {"parts": narrate.fold_parts(parts, n_states), "clicks": clicks, "builds": b}

    def test_matching(self):
        steps, warnings = narrate.plan_slide(self.slide("A [click] B", 2), 1.0)
        self.assertEqual(warnings, [])
        self.assertEqual([len(s["parts"]) for s in steps], [1, 1])

    def test_too_few_clicks_warns(self):
        _, warnings = narrate.plan_slide(self.slide("A [click] B", 4), 1.0)
        self.assertIn("play after the narration", warnings[0])

    def test_too_many_clicks_warns(self):
        _, warnings = narrate.plan_slide(self.slide("A [click] B [click] C", 2), 1.0)
        self.assertIn("stays on the last build", warnings[0])


if __name__ == "__main__":
    unittest.main()
