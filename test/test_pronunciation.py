#!/usr/bin/env python3
"""Unit tests for --check-pronunciation, --fix-word and the --sample cache. The checks
use Kokoro's phonemizer (misaki + spaCy, no voice model); the sample tests mock the TTS."""
import contextlib
import io
import os
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import narrate  # noqa: E402


def make_deck(path, notes_runs):
    """A deck with one slide per entry; each entry is a list of note runs (one paragraph)."""
    from pptx import Presentation
    prs = Presentation()
    for runs in notes_runs:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        para = slide.notes_slide.notes_text_frame.paragraphs[0]
        for text in runs:
            para.add_run().text = text
    prs.save(str(path))


def notes(path, n):
    from pptx import Presentation
    return [r.text for r in Presentation(str(path)).slides[n - 1].notes_slide.notes_text_frame.paragraphs[0].runs]


def quiet():
    return contextlib.redirect_stderr(io.StringIO())


class Spans(unittest.TestCase):
    def test_nth_word_skips_stage_directions_but_counts_overrides(self):
        t = "Read [read slowly] the [read](/ɹˈɛd/) and read-only notes; already read."
        self.assertEqual(t[slice(*narrate.nth_word_span(t, "read", 1))], "Read")
        self.assertEqual(narrate.nth_word_span(t, "read", 2)[0], t.index("[read]") + 1)
        self.assertEqual(narrate.nth_word_span(t, "read", 3)[0], t.index("read-only"))
        self.assertEqual(narrate.nth_word_span(t, "read", 4)[0], t.rindex("read"))
        self.assertIsNone(narrate.nth_word_span(t, "read", 5))
        self.assertIsNone(narrate.nth_word_span("bread and reader", "read", 1))

    def test_with_override_adds_or_replaces(self):
        self.assertEqual(narrate.with_override("I read it.", "read", 1, "ɹˈɛd"), "I [read](/ɹˈɛd/) it.")
        self.assertEqual(narrate.with_override("I [read](/ɹˈid/) it.", "read", 1, "ɹˈɛd"), "I [read](/ɹˈɛd/) it.")
        self.assertEqual(narrate.with_override("Read it.", "read", 1, "ɹˈid"), "[Read](/ɹˈid/) it.")


class FixWord(unittest.TestCase):
    def setUp(self):
        self.dir = Path(tempfile.mkdtemp())
        self.deck = self.dir / "talk.pptx"
        make_deck(self.deck, [["I read the record. ", "Then we ", "record", " it."], ["Nothing."]])

    def test_fixes_the_right_occurrence_and_keeps_runs(self):
        with quiet():
            narrate.fix_word(self.deck, 1, "record", 2, "/ɹəkˈɔɹd/")
        self.assertEqual(notes(self.deck, 1), ["I read the record. ", "Then we ", "[record](/ɹəkˈɔɹd/)", " it."])
        backup = self.dir / "talk (before deck-narrator).pptx"
        self.assertTrue(backup.is_file())
        self.assertEqual(notes(backup, 1)[2], "record")
        with quiet():                                   # again: replaces the override; backup stays the original
            narrate.fix_word(self.deck, 1, "record", 2, "ɹˈɛkəɹd")
            narrate.fix_word(self.deck, 1, "read", 1, "ɹˈɛd")
        self.assertEqual(notes(self.deck, 1), ["I [read](/ɹˈɛd/) the record. ", "Then we ", "[record](/ɹˈɛkəɹd/)", " it."])
        self.assertEqual(notes(backup, 1)[2], "record")

    def test_bad_requests_fail_cleanly(self):
        for args in ((1, "record", 3, "x"), (2, "record", 1, "x"), (9, "read", 1, "x"), (1, "read", 1, "a/b")):
            with self.subTest(args=args), self.assertRaises(narrate.Fail):
                narrate.fix_word(self.deck, *args)


class Issues(unittest.TestCase):
    def slides(self, text, voice="af_heart"):
        return [{"number": 1, "parts": [{"voice": voice, "text": text, "step": 0}]}]

    def test_finds_heteronyms_unknown_and_spelled_words(self):
        issues = narrate.pronunciation_issues(self.slides(
            "Yesterday I read the record. We record a demo, then OPSWAT will lead. Zscaler can read it too."))
        found = [(i["word"], i["n"], i["kind"]) for i in issues]
        self.assertIn(("read", 1, "heteronym"), found)
        self.assertIn(("record", 2, "heteronym"), found)
        self.assertIn(("read", 2, "heteronym"), found)
        self.assertIn(("lead", 1, "heteronym"), found)          # from EXTRA_HETERONYMS
        self.assertIn(("OPSWAT", 1, "spelled"), found)
        self.assertIn(("Zscaler", 1, "unknown"), found)
        read2 = next(i for i in issues if i["word"] == "read" and i["n"] == 2)
        self.assertEqual(read2["sentence"], "Zscaler can read it too.")
        self.assertTrue(all("[read](/" in o["say"] for o in read2["options"]))
        self.assertGreaterEqual(len(read2["options"]), 2)

    def test_overrides_and_plain_words_are_left_alone(self):
        issues = narrate.pronunciation_issues(self.slides("I [read](/ɹˈɛd/) the notes. The cat sat on the mat."))
        self.assertEqual(issues, [])


class SampleCache(unittest.TestCase):
    def setUp(self):
        self.dir = Path(tempfile.mkdtemp())
        self.env = mock.patch.dict(os.environ, {"DECK_NARRATOR_SAMPLE_CACHE": str(self.dir / "cache"),
                                                "DECK_NARRATOR_VOICE_DIR": str(self.dir / "voices")})
        self.env.start()
        self.addCleanup(self.env.stop)

    def run_sample(self, tts, out):
        args = SimpleNamespace(output=str(out), speed=1.0, pronunciations=None, sample="Hello there.",
                               voice="af_heart", device="cpu")
        with mock.patch.object(narrate, "TTS", tts), mock.patch.object(narrate, "app_dir", return_value=self.dir), \
                quiet(), contextlib.redirect_stdout(io.StringIO()):
            narrate.speak_sample(args)

    def test_second_play_comes_from_the_cache(self):
        import numpy as np
        tts = mock.Mock()
        tts.return_value.synthesize.return_value = np.zeros(2400, dtype=np.float32)
        self.run_sample(tts, self.dir / "a.wav")
        self.run_sample(tts, self.dir / "b.wav")
        self.assertEqual(tts.return_value.synthesize.call_count, 1)
        self.assertEqual((self.dir / "a.wav").read_bytes(), (self.dir / "b.wav").read_bytes())

    def test_shipped_samples_are_used_first(self):
        shipped = self.dir / "samples"
        shipped.mkdir()
        (shipped / (narrate.sample_key("af_heart", "Hello there.", 1.0) + ".wav")).write_bytes(b"RIFF-shipped")
        tts = mock.Mock(side_effect=AssertionError("rendered"))
        self.run_sample(tts, self.dir / "c.wav")
        self.assertEqual((self.dir / "c.wav").read_bytes(), b"RIFF-shipped")

    def test_key_changes_with_what_you_hear(self):
        k = narrate.sample_key("af_heart", "Hi.", 1.0)
        self.assertNotEqual(k, narrate.sample_key("af_bella", "Hi.", 1.0))
        self.assertNotEqual(k, narrate.sample_key("af_heart", "Hi!", 1.0))
        self.assertNotEqual(k, narrate.sample_key("af_heart", "Hi.", 1.1))
        (self.dir / "voices").mkdir()
        (self.dir / "voices" / "af_heart.pt").write_bytes(b"updated voice")
        self.assertNotEqual(k, narrate.sample_key("af_heart", "Hi.", 1.0))   # an updated voice


if __name__ == "__main__":
    unittest.main()
